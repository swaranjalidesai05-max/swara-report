import sys
sys.path.insert(0, ".")
from app import *
import json

init_db()
conn = get_db()
event = conn.execute("SELECT * FROM events ORDER BY id DESC LIMIT 1").fetchone()
conn.close()
event = dict(event)
print("Event: " + str(event.get("title")))

create_default_template()
from docx import Document
doc = Document(TEMPLATE_PATH)

funcs = [
    ("replace_placeholders", lambda: replace_placeholders(doc, {"{{event_name}}": event.get("title", "")})),
    ("insert_event_details", lambda: insert_event_details_paragraph(doc, "<<EVENT_DETAILS>>", event)),
    ("insert_full_page_image (permission)", lambda: insert_full_page_image(doc, "<<IMAGE_PERMISSION>>", event.get("permission_letter"), heading="Permission Letter")),
    ("insert_full_page_image (invitation)", lambda: insert_full_page_image(doc, "<<IMAGE_INVITATION>>", event.get("invitation_letter"), heading="Invitation Letter")),
    ("insert_full_page_image (notice)", lambda: insert_full_page_image(doc, "<<IMAGE_NOTICE>>", event.get("notice_letter"), heading="Notice")),
    ("insert_full_page_image (appreciation)", lambda: insert_full_page_image(doc, "<<IMAGE_APPRECIATION>>", event.get("appreciation_letter"), heading="Appreciation Letter")),
    ("insert_event_photos", lambda: insert_event_photos(doc, "<<EVENT_PHOTOS>>", json.loads(event["event_photos"]) if event.get("event_photos") else [])),
    ("insert_attendance", lambda: insert_attendance(doc, "<<ATTENDANCE_FILE>>", event.get("attendance_photo"))),
    ("insert_feedback_table", lambda: insert_feedback_table(doc, "<<FEEDBACK_TABLE>>", json.loads(event["feedback_data"]) if event.get("feedback_data") else [])),
]

for name, fn in funcs:
    try:
        fn()
        print(name + ": OK")
    except Exception as e:
        print(name + " ERROR: " + str(e))

doc.save("generated_reports/test_noempty.docx")
print("ALL DONE")
